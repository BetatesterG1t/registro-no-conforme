"""Regenera data/validaciones.json y data/ordenes.json desde los Excel de PNCIPROY2."""
from pathlib import Path
import json
from openpyxl import load_workbook

DESKTOP = Path.home() / "Desktop"
CANDIDATOS = [
    DESKTOP / "Registro No Conforme - pagina corregida",
    DESKTOP / "PROYECTO" / "PNCIPROY2",
    DESKTOP / "PROYECTO" / "PNCIPROY",
    DESKTOP / "PNCIPROY",
]


def primer_existente(nombre):
    for carpeta in CANDIDATOS:
        ruta = carpeta / nombre
        if ruta.exists():
            return ruta
    raise FileNotFoundError("No se encontró " + nombre + " en: " + ", ".join(str(c) for c in CANDIDATOS))


VALIDACIONES_XLSX = primer_existente("VALIDACIONES.xlsx")
SEGUIMIENTO_XLSX = primer_existente("SEGUIMIENTO DE OT1 (1).xlsx")
OUT = Path(__file__).resolve().parents[1] / "data"


def col_values(ws, col_idx):
    seen = set()
    vals = []
    for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx, values_only=True):
        v = row[0]
        if v is None:
            continue
        s = str(v).strip()
        if s and s not in seen:
            seen.add(s)
            vals.append(s)
    return vals


def extraer_validaciones():
    wb = load_workbook(VALIDACIONES_XLSX, data_only=True)
    ws = wb.active
    data = {
        "mes": col_values(ws, 1),
        "proceso_detecta": col_values(ws, 4),
        "maquina_detecta": col_values(ws, 7),
        "operador_detecta": col_values(ws, 10),
        "proceso_origina": col_values(ws, 13),
        "maquina": col_values(ws, 16),
        "operador": col_values(ws, 19),
        "defecto": col_values(ws, 22),
        "reporta": col_values(ws, 25),
        "autorizo": col_values(ws, 28),
    }
    wb.close()
    return data


def extraer_ordenes():
    wb = load_workbook(SEGUIMIENTO_XLSX, read_only=True, data_only=True)
    ws = wb["rep-ov-espec"]
    registros = []
    vacias = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        orden, oc, cliente = row[1], row[2], row[5]
        sku, id_art, producto = row[10], row[11], row[12]
        if (orden is None or str(orden).strip() == "") and (oc is None or str(oc).strip() == ""):
            vacias += 1
            if vacias >= 20:
                break
            continue
        vacias = 0
        registros.append({
            "orden": str(orden).strip() if orden is not None else "",
            "ordenCompra": str(oc).strip() if oc is not None else "",
            "cliente": str(cliente).strip() if cliente is not None else "",
            "sku": str(sku).strip() if sku is not None else "",
            "idArticulo": str(id_art).strip() if id_art is not None else "",
            "producto": str(producto).strip() if producto is not None else "",
        })
    wb.close()
    return registros


def main():
    print("VALIDACIONES:", VALIDACIONES_XLSX)
    print("SEGUIMIENTO:", SEGUIMIENTO_XLSX)
    OUT.mkdir(parents=True, exist_ok=True)
    validaciones = extraer_validaciones()
    (OUT / "validaciones.json").write_text(
        json.dumps(validaciones, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    ordenes = extraer_ordenes()
    (OUT / "ordenes.json").write_text(
        json.dumps(ordenes, ensure_ascii=False), encoding="utf-8"
    )
    print("validaciones.json:", {k: len(v) for k, v in validaciones.items()})
    print("ordenes.json:", len(ordenes), "registros")


if __name__ == "__main__":
    main()
