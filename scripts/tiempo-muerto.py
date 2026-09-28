"""Arma data/tiempo-muerto.json y lo incrusta en tiempo-muerto.html.

Lee los libros de tiempo muerto del área de impresión (uno por mes).
No copia los Excel al repositorio: solo escribe JSON y HTML.

Uso:
  python scripts/tiempo-muerto.py /ruta/con/los/xlsm
"""
from __future__ import annotations

import json
import sys
import unicodedata
from collections import defaultdict
from datetime import datetime, time, timedelta
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
HTML_PATH = ROOT / "tiempo-muerto.html"
JSON_PATH = ROOT / "data" / "tiempo-muerto.json"

MESES = [
    ("2026-01", "Enero", "Ene", ("ENERO",)),
    ("2026-02", "Febrero", "Feb", ("FEBRERO",)),
    ("2026-03", "Marzo", "Mar", ("MARZO",)),
    ("2026-04", "Abril", "Abr", ("ABRIL",)),
    ("2026-05", "Mayo", "May", ("MAYO",)),
    ("2026-06", "Junio", "Jun", ("JUNIO",)),
    ("2026-07", "Julio", "Jul", ("JULIO",)),
    ("2026-08", "Agosto", "Ago", ("AGOSTO",)),
    ("2026-09", "Septiembre", "Sep", ("SEPTIEMBRE",)),
]

CONCEPTOS = [
    ("vobo", "Visto bueno (VoBo)", "aprobacion", True),
    ("tintas", "Tintas", "materiales", True),
    ("maquina", "Máquina descompuesta, en reparación", "maquina", True),
    ("calidad", "Autorización de calidad", "aprobacion", True),
    ("mangas", "Mangas y/o anilox en mal estado", "materiales", True),
    ("secuencia", "Cambio de secuencia", "programa", True),
    ("espera_montaje", "Espera de montaje de grabados", "montaje", True),
    ("grabados", "Grabados", "montaje", True),
    ("falta_material", "Falta de materiales", "materiales", True),
    ("preprensa", "Preprensa", "montaje", True),
    ("montaje_mangas", "Montaje por mismas mangas", "montaje", True),
    ("orden_corta", "Cambios excesivos por orden corta", "programa", True),
    ("material_mal", "Material en mal estado", "materiales", True),
    ("programacion", "Cambio de programación", "programa", True),
    ("luz", "Luz", "maquina", True),
    ("error_montaje", "Error en montaje", "montaje", True),
    ("orden_suspendida", "Orden suspendida", "programa", True),
    ("alarmada", "Máquina alarmada", "maquina", False),
]

GRUPOS = [
    ("aprobacion", "Aprobación", "Visto bueno y liberación de calidad."),
    ("materiales", "Materiales", "Tintas, sustrato, mangas y anilox."),
    ("maquina", "Máquina", "Equipo detenido, alarma o falta de luz."),
    ("montaje", "Montaje", "Preprensa, grabados y montaje."),
    ("programa", "Programa", "Cambios de orden y órdenes suspendidas."),
]

# Orden específico: frases largas antes que las que contienen.
REGLAS = [
    ("alarmada", "alarmada"),
    ("visto bueno", "vobo"),
    ("vobo", "vobo"),
    ("tintas", "tintas"),
    ("descompuesta", "maquina"),
    ("espera de montaje", "espera_montaje"),
    ("mismas mangas", "montaje_mangas"),
    ("error en montaje", "error_montaje"),
    ("anilox", "mangas"),
    ("secuencia", "secuencia"),
    ("falta de material", "falta_material"),
    ("pre prensa", "preprensa"),
    ("preprensa", "preprensa"),
    ("orden corta", "orden_corta"),
    ("calidad", "calidad"),
    ("mal estado", "material_mal"),
    ("programacion", "programacion"),
    ("luz", "luz"),
    ("suspendida", "orden_suspendida"),
    ("grabados", "grabados"),
]

MES_CORTO = ["ene", "feb", "mar", "abr", "may", "jun", "jul", "ago", "sep", "oct", "nov", "dic"]
ORDEN_MAQUINAS = ["F4-1", "F4-2", "F4-3", "F2-1", "F2-2", "F2-3", "FJ-2", "FLEXO", "FW"]


def fold(value) -> str:
    text = str(value).replace(":", " ")
    text = "".join(
        ch for ch in unicodedata.normalize("NFD", text) if unicodedata.category(ch) != "Mn"
    )
    return " ".join(text.lower().split())


def canon(label: str):
    folded = fold(label)
    if folded.startswith("total") or folded.startswith("tiempo muerto"):
        return None
    for needle, cid in REGLAS:
        if needle in folded:
            return cid
    return None


def horas_celda(value) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, timedelta):
        return value.total_seconds() / 3600.0
    if isinstance(value, time):
        return value.hour + value.minute / 60.0 + value.second / 3600.0
    if isinstance(value, datetime):
        return value.hour + value.minute / 60.0 + value.second / 3600.0
    if isinstance(value, (int, float)):
        return float(value)
    return 0.0


def redondear(value, ndigits=4):
    return round(float(value) + 1e-9, ndigits)


def norm_maquina(value: str) -> str:
    name = " ".join(str(value).strip().split()).upper()
    return name.replace("FJ2", "FJ-2")


def buscar_libros(carpeta: Path) -> dict:
    encontrados = {}
    for path in sorted(carpeta.glob("*.xlsm")):
        nombre = path.name.upper()
        for mes_id, _nombre, _corto, claves in MESES:
            if any(clave in nombre for clave in claves) and mes_id not in encontrados:
                encontrados[mes_id] = path
                break
    faltan = [mes_id for mes_id, *_ in MESES if mes_id not in encontrados]
    if faltan:
        raise SystemExit(f"Faltan libros de: {', '.join(faltan)} en {carpeta}")
    return encontrados


def hoja(wb, prefijo: str):
    for name in wb.sheetnames:
        if name.upper().startswith(prefijo):
            return name
    raise KeyError(prefijo)


def filas(path: Path, prefijo: str):
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[hoja(wb, prefijo)]
    data = [list(row) for row in ws.iter_rows(values_only=True)]
    wb.close()
    return data


def parse_grafico(rows):
    conceptos = {}
    semanas_concepto = {}
    desconocidos = []
    for row in rows:
        label = row[0] if row else None
        if not isinstance(label, str) or not label.strip():
            continue
        if fold(label).startswith("concepto"):
            continue
        cid = canon(label)
        if cid is None:
            if not fold(label).startswith("tiempo muerto"):
                desconocidos.append(label)
            continue
        weeks = []
        for col in range(1, 6):
            value = row[col] if col < len(row) else None
            weeks.append(None if value is None else float(value))
        total = row[6] if len(row) > 6 and isinstance(row[6], (int, float)) else None
        if total is None:
            total = sum(w or 0 for w in weeks)
        conceptos[cid] = float(total)
        semanas_concepto[cid] = weeks

    total_row = None
    for index, row in enumerate(rows):
        marker = row[1] if len(row) > 1 else None
        if isinstance(marker, str) and "TM" in marker.upper() and index + 1 < len(rows):
            total_row = rows[index + 1]
            break
    if total_row is None:
        raise SystemExit("No encontré la fila de total del gráfico mensual")

    semanas = []
    for col in range(1, 6):
        value = total_row[col] if col < len(total_row) else None
        semanas.append(None if value is None else float(value))
    for col in range(5):
        if all((semanas_concepto[cid][col] is None) for cid in semanas_concepto):
            semanas[col] = None
    activas = sum(1 for value in semanas if value is not None)

    def num(col):
        value = total_row[col] if col < len(total_row) else 0
        return float(value or 0)

    return {
        "conceptos": conceptos,
        "semanas_concepto": semanas_concepto,
        "semanas": semanas,
        "semanas_activas": activas,
        "total_horas": num(6),
        "ml": num(7),
        "kg": num(8),
        "costo_teorico": num(9),
        "costo_facturacion": num(10),
        "desconocidos": desconocidos,
    }


def horas_fila(row, col):
    if col + 15 >= len(row):
        return 0.0
    cell = row[col + 15]
    if cell is None:
        total = 0.0
        for offset in (5, 8, 11, 14):
            if col + offset < len(row):
                total += horas_celda(row[col + offset])
        return total
    return horas_celda(cell)


def parse_semanal(rows):
    maquinas = []
    if not rows:
        return {}, {}, [], []
    for index, value in enumerate(rows[0]):
        if isinstance(value, str) and value.strip():
            maquinas.append((index, norm_maquina(value)))
    dias = [
        index
        for index, row in enumerate(rows)
        if row and isinstance(row[0], datetime) and row[0].year > 2000
    ]
    nombres = {nombre.casefold() for _, nombre in maquinas}
    por_maquina = defaultdict(float)
    por_maquina_concepto = defaultdict(lambda: defaultdict(float))
    desconocidos = []
    for pos, inicio in enumerate(dias):
        fin = dias[pos + 1] if pos + 1 < len(dias) else len(rows)
        actual = None
        for row_index in range(inicio + 1, fin):
            row = rows[row_index]
            label = row[0] if row else None
            if isinstance(label, str) and label.strip():
                if label.strip().casefold() in nombres:
                    actual = None
                    continue
                cid = canon(label)
                actual = cid
                if cid is None and not fold(label).startswith("total") and not fold(label).startswith("tiempo muerto"):
                    desconocidos.append(label.strip())
            if not actual:
                continue
            for col, nombre in maquinas:
                horas = horas_fila(row, col)
                if horas > 0.0000001:
                    por_maquina[nombre] += horas
                    por_maquina_concepto[nombre][actual] += horas
    fechas = [rows[index][0].date() for index in dias]
    nombres_orden = []
    for _, nombre in maquinas:
        if nombre not in nombres_orden:
            nombres_orden.append(nombre)
    return por_maquina, por_maquina_concepto, fechas, sorted(set(desconocidos)), nombres_orden


def rango_fechas(fechas) -> str:
    if not fechas:
        return ""
    inicio, fin = fechas[0], fechas[-1]
    if inicio.month == fin.month:
        return f"{inicio.day}–{fin.day} {MES_CORTO[inicio.month - 1]}"
    return (
        f"{inicio.day} {MES_CORTO[inicio.month - 1]} – "
        f"{fin.day} {MES_CORTO[fin.month - 1]}"
    )


def cuadrar_semana(grafico, por_maquina_concepto):
    detalle = defaultdict(float)
    for conceptos in por_maquina_concepto.values():
        for cid, horas in conceptos.items():
            detalle[cid] += horas
    mejor = None
    for indice, total in enumerate(grafico["semanas"]):
        if total is None:
            continue
        diferencia = 0.0
        ids = set(grafico["semanas_concepto"]) | set(detalle)
        for cid in ids:
            if cid == "alarmada":
                continue
            semana = grafico["semanas_concepto"].get(cid, [None] * 5)[indice]
            diferencia += abs((detalle.get(cid, 0.0)) - (0.0 if semana is None else semana))
        if mejor is None or diferencia < mejor[0]:
            mejor = (diferencia, indice, total)
    return detalle, mejor


def limpiar(mapa):
    return {llave: redondear(valor) for llave, valor in mapa.items() if valor > 0.0000001}


def construir(carpeta: Path):
    libros = buscar_libros(carpeta)
    meses = []
    maquinas = []
    advertencias = []
    for mes_id, nombre, corto, _claves in MESES:
        path = libros[mes_id]
        grafico = parse_grafico(filas(path, "GRAFICO"))
        por_maquina, por_mc, fechas, desconocidos, nombres_mes = parse_semanal(filas(path, "TM SEMANAL"))
        for extra in grafico["desconocidos"] + desconocidos:
            advertencias.append(f"{nombre}: concepto no reconocido «{extra}»")
        for nombre_maq in nombres_mes:
            if nombre_maq not in maquinas:
                maquinas.append(nombre_maq)
        detalle, mejor = cuadrar_semana(grafico, por_mc)
        if mejor is None:
            raise SystemExit(f"{nombre}: no hay semana activa en el gráfico")
        diferencia, indice, total_semana = mejor
        alarmada = detalle.get("alarmada", 0.0)
        detalle_oficial = sum(valor for cid, valor in detalle.items() if cid != "alarmada")
        mes_archivo = int(mes_id.split("-")[1])
        confiables = bool(fechas) and any(fecha.month == mes_archivo for fecha in fechas)
        if diferencia > 0.08:
            advertencias.append(
                f"{nombre}: la semana {indice + 1} difiere {diferencia:.2f} h del detalle"
            )
        meses.append(
            {
                "id": mes_id,
                "nombre": nombre,
                "corto": corto,
                "parcial": mes_id == "2026-09",
                "archivo": path.name,
                "semanas": [None if value is None else redondear(value, 4) for value in grafico["semanas"]],
                "semanas_activas": grafico["semanas_activas"],
                "total_horas": redondear(grafico["total_horas"], 4),
                "ml": redondear(grafico["ml"], 2),
                "kg": redondear(grafico["kg"], 2),
                "costo_teorico": redondear(grafico["costo_teorico"], 2),
                "costo_facturacion": redondear(grafico["costo_facturacion"], 2),
                "conceptos": {cid: redondear(valor) for cid, valor in grafico["conceptos"].items()},
                "semanas_concepto": {
                    cid: [None if value is None else redondear(value) for value in weeks]
                    for cid, weeks in grafico["semanas_concepto"].items()
                },
                "detalle": {
                    "semana_num": indice + 1,
                    "etiqueta": f"Semana {indice + 1}",
                    "rango": rango_fechas(fechas) if confiables else "",
                    "fechas_confiables": confiables,
                    "horas_oficiales_semana": redondear(total_semana, 4),
                    "horas_detalle": redondear(sum(detalle.values()), 4),
                    "horas_sin_alarmada": redondear(detalle_oficial, 4),
                    "diferencia": redondear(detalle_oficial - total_semana, 4),
                    "alarmada": redondear(alarmada, 4),
                    "por_maquina": limpiar(por_maquina),
                    "por_maquina_concepto": {
                        nombre_maq: limpiar(conceptos) for nombre_maq, conceptos in por_mc.items()
                    },
                },
            }
        )

    ordenadas = [nombre for nombre in ORDEN_MAQUINAS if nombre in maquinas]
    ordenadas += [nombre for nombre in maquinas if nombre not in ordenadas]
    factor_fact = 156250 / 12
    factor_teo = 13081.6
    data = {
        "meta": {
            "titulo": "Tiempo muerto",
            "area": "Área de impresión",
            "anio": 2026,
            "corte": "20 de septiembre de 2026",
            "periodo": "Enero al 20 de septiembre de 2026",
            "factores": {
                "ml_por_hora": 8760,
                "kg_por_hora": 116.8,
                "costo_teorico_por_hora": factor_teo,
                "costo_facturacion_por_hora": factor_fact,
            },
            "advertencias": advertencias,
        },
        "grupos": [
            {"id": gid, "nombre": nombre, "texto": texto}
            for gid, nombre, texto in GRUPOS
        ],
        "conceptos": [
            {"id": cid, "nombre": nombre, "grupo": grupo, "en_grafico": en_grafico}
            for cid, nombre, grupo, en_grafico in CONCEPTOS
        ],
        "maquinas": ordenadas,
        "meses": meses,
    }
    validar(data)
    return data


def validar(data):
    enero = data["meses"][0]
    if abs(enero["total_horas"] - 363.44) > 0.02:
        raise SystemExit(f"Total de enero inesperado: {enero['total_horas']}")
    if abs(enero["detalle"]["horas_oficiales_semana"] - 104.25) > 0.02:
        raise SystemExit("La semana guardada de enero no cuadra con 104.25 h")
    for mes in data["meses"]:
        suma = sum(mes["conceptos"].values())
        if abs(suma - mes["total_horas"]) > 0.15:
            raise SystemExit(
                f"{mes['nombre']}: conceptos {suma:.2f} vs total {mes['total_horas']:.2f}"
            )
        if abs(mes["detalle"]["diferencia"]) > 0.08:
            raise SystemExit(
                f"{mes['nombre']}: detalle {mes['detalle']['horas_sin_alarmada']:.2f} "
                f"vs semana {mes['detalle']['horas_oficiales_semana']:.2f}"
            )
    print("Conciliación correcta.")
    total = sum(mes["total_horas"] for mes in data["meses"])
    print(f"Horas oficiales del año: {total:.2f}")
    for mes in data["meses"]:
        det = mes["detalle"]
        print(
            f"  {mes['corto']:3} {mes['total_horas']:8.2f} h  "
            f"semana {det['semana_num']} {det['horas_oficiales_semana']:7.2f}  "
            f"detalle {det['horas_detalle']:7.2f}  alarmada {det['alarmada']:5.2f}  "
            f"fechas {'ok' if det['fechas_confiables'] else 'no'} {det['rango']}"
        )


def incrustar_html(data):
    if not HTML_PATH.exists():
        return
    html = HTML_PATH.read_text(encoding="utf-8")
    inicio = html.find('<script id="tm-data" type="application/json">')
    fin = html.find("</script>", inicio)
    if inicio < 0 or fin < 0:
        raise SystemExit("tiempo-muerto.html no tiene el bloque tm-data")
    payload = json.dumps(data, ensure_ascii=False, separators=(",", ":")).replace("<", "\\u003c")
    nuevo = (
        html[:inicio]
        + '<script id="tm-data" type="application/json">\n'
        + payload
        + "\n"
        + html[fin:]
    )
    HTML_PATH.write_text(nuevo, encoding="utf-8")


def main():
    if len(sys.argv) < 2:
        raise SystemExit("Indica la carpeta con los .xlsm")
    carpeta = Path(sys.argv[1]).expanduser().resolve()
    data = construir(carpeta)
    JSON_PATH.parent.mkdir(parents=True, exist_ok=True)
    JSON_PATH.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    incrustar_html(data)
    print(f"JSON: {JSON_PATH}")
    print(f"HTML: {HTML_PATH}")


if __name__ == "__main__":
    main()
